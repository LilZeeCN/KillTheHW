#!/usr/bin/env python3
"""Inspect a DOCX template and fill explicit, hash-bound regions without rebuilding it."""
import argparse
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile

from _common import digest, write_json


def index_document(document):
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.text.paragraph import Paragraph
    stories = [('document', document.part)]
    stories += [(part.partname, part) for rel in document.part.rels.values()
                if rel.reltype in {RT.HEADER, RT.FOOTER} and not rel.is_external
                for part in [rel.target_part]]
    nodes, entries, seen = {}, [], set()
    counts = {'P': 0, 'CELL': 0}
    for n, node in enumerate(document._element.body, 1):
        if node.tag == qn('w:sectPr'):
            continue
        key = f'B{n:04d}'
        nodes[key] = node
        entries.append({'id': key, 'kind': node.tag.split('}')[-1],
                        'text': ''.join(t.text or '' for t in node.iter(qn('w:t'))), 'story': 'document'})
    for story, part in stories:
        if id(part) in seen:
            continue
        seen.add(id(part))
        for node in part.element.iter():
            prefix = 'P' if node.tag == qn('w:p') else 'CELL' if node.tag == qn('w:tc') else None
            if not prefix:
                continue
            counts[prefix] += 1
            key = f'{prefix}{counts[prefix]:04d}'
            nodes[key] = node
            text = Paragraph(node, document).text if prefix == 'P' else '\n'.join(
                Paragraph(p, document).text for p in node.iter(qn('w:p')))
            entry = {'id': key, 'story': str(story), 'text': text,
                     'images': len(list(node.iter(qn('w:drawing')))),
                     'xml_path': node.getroottree().getpath(node)}
            if prefix == 'P':
                st = Paragraph(node, document).style
                entry['style'] = st.name if st is not None else 'Normal'
            nodes[key] = node
            entries.append(entry)
    return nodes, entries, stories


def apply_font_formatting(run, font_ascii=None, font_eastasia=None, font_size_pt=None, bold=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    if font_ascii or font_eastasia:
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        if font_ascii:
            run.font.name = font_ascii
            rFonts.set(qn('w:ascii'), font_ascii)
            rFonts.set(qn('w:hAnsi'), font_ascii)
        if font_eastasia:
            rFonts.set(qn('w:eastAsia'), font_eastasia)
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph(node, document, text):
    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(node, document)
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)
    return paragraph


def replace_token(paragraph, token, value):
    from docx.oxml.ns import qn
    runs = paragraph.runs
    text = ''.join(run.text for run in runs)
    starts, offset = [], 0
    for run in runs:
        starts.append(offset)
        offset += len(run.text)
    positions, pos = [], text.find(token)
    while pos >= 0:
        positions.append(pos)
        pos = text.find(token, pos + len(token))
    for pos in reversed(positions):
        selected = [i for i, run in enumerate(runs) if starts[i] < pos + len(token) and starts[i] + len(run.text) > pos]
        for i in selected:
            if any(list(runs[i]._r.iter(qn(tag))) for tag in ('w:drawing', 'w:fldChar', 'w:instrText')):
                raise ValueError('Token shares a run with a drawing/field; use an explicit field region')
        first, last = selected[0], selected[-1]
        before = runs[first].text[:pos - starts[first]]
        after = runs[last].text[pos + len(token) - starts[last]:]
        runs[first].text = before + value + (after if first == last else '')
        for i in selected[1:]:
            runs[i].text = after if i == last else ''
    return len(positions)


def build(template, plan_path, output):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.shared import Cm, Pt
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    if digest(template) != plan['template_sha256']:
        raise ValueError('Template changed; inspect again and revise the plan')
    if output.exists() or output.resolve() == template.resolve():
        raise ValueError('Output must be a new file, separate from the template')
    document = Document(template)
    nodes, entries, stories = index_document(document)
    ref_doc = Document((plan_path.parent / plan['reference_docx']).resolve()) if plan.get('reference_docx') else None
    ref_nodes = index_document(ref_doc)[0] if ref_doc else {}
    body = document._element.body
    children = list(body)
    regions, claimed = [], set()
    # Resolve all original anchors before any editing; duplicated headings are harmless.
    for region in plan.get('sections', []):
        start = nodes[region['start']]
        end = nodes[region['end']] if region.get('end') != 'END' else body[-1]
        if start.getparent() is not body or end.getparent() is not body or start.tag != qn('w:p'):
            raise ValueError('Section anchors must be direct body blocks, with a paragraph start')
        left, right = children.index(start), children.index(end)
        if right <= left or claimed.intersection(range(left, right)):
            raise ValueError('Section ranges overlap or are reversed')
        claimed.update(range(left, right))
        removed = children[left + 1:right]
        if any(list(node.iter(qn('w:sectPr'))) for node in removed):
            raise ValueError('Range contains section properties; preserve the section boundary explicitly')
        if region.get('expect') is not None and Paragraph(start, document).text != region['expect']:
            raise ValueError(f"Section heading mismatch: {region['start']}")
        regions.append((start, end, removed, region))
    for key in plan.get('remove_blocks', []):
        node = nodes[key]
        if (node.getparent() is not body or children.index(node) in claimed
                or any(node is end for _, end, _, _ in regions)):
            raise ValueError(f'Removal conflicts with a section or is not a body block: {key}')
        if list(node.iter(qn('w:sectPr'))):
            raise ValueError('Cannot remove a section boundary')
        body.remove(node)
    for replacement in plan.get('replacements', []):
        token = replacement['token']
        if not token:
            raise ValueError('Empty replacement token')
        count = sum(replace_token(Paragraph(node, document), token, str(replacement['value']))
                    for key, node in nodes.items() if key.startswith('P') and node.getparent() is not None)
        if count != replacement['expected_count']:
            raise ValueError(f'Token count mismatch for {token!r}: expected {replacement["expected_count"]}, found {count}')
    for field in plan.get('fields', []):
        node = nodes[field['target']]
        if node.tag == qn('w:tc'):
            paragraphs = list(node.iter(qn('w:p')))
            actual = '\n'.join(Paragraph(p, document).text for p in paragraphs)
        elif node.tag == qn('w:p'):
            paragraphs, actual = [node], Paragraph(node, document).text
        else:
            raise ValueError('Field target must be a paragraph or cell')
        if actual != field['expect']:
            raise ValueError(f"Field content mismatch: {field['target']}")
        p = node.find(qn('w:p')) if node.tag == qn('w:tc') else node
        if p is None:
            from docx.oxml import OxmlElement
            p = OxmlElement('w:p')
            node.append(p)
        if node.tag == qn('w:tc'):
            for child in list(node):
                if child is not p and child.tag != qn('w:tcPr'):
                    node.remove(child)
        set_paragraph(p, document, str(field['text']))

    def paragraph_for(block):
        paragraph = document.add_paragraph()
        if block.get('style_from'):
            sf = block['style_from']
            original = Paragraph(ref_nodes[sf[4:]], ref_doc) if sf.startswith('REF:') else Paragraph(nodes[sf], document)
            if original._p.pPr is not None:
                if paragraph._p.pPr is not None:
                    paragraph._p.remove(paragraph._p.pPr)
                paragraph._p.insert(0, deepcopy(original._p.pPr))
            if original.runs and original.runs[0]._r.rPr is not None:
                paragraph.add_run()._r.insert(0, deepcopy(original.runs[0]._r.rPr))
        if block.get('style'):
            paragraph.style = block['style']
        if block.get('line_spacing'):
            paragraph.paragraph_format.line_spacing = float(block['line_spacing'])
        if block.get('alignment'):
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            align_map = {
                'left': WD_ALIGN_PARAGRAPH.LEFT,
                'center': WD_ALIGN_PARAGRAPH.CENTER,
                'right': WD_ALIGN_PARAGRAPH.RIGHT,
                'justify': WD_ALIGN_PARAGRAPH.JUSTIFY
            }
            if block['alignment'] in align_map:
                paragraph.alignment = align_map[block['alignment']]
        return paragraph

    for start, end, removed, region in reversed(regions):
        for node in removed:
            body.remove(node)
        if 'heading' in region:
            set_paragraph(start, document, region['heading'])
        for block in region['blocks']:
            kind = block['type']
            if kind not in {'paragraph', 'code', 'image'}:
                raise ValueError(f'Unknown block type: {kind}')
            paragraph = paragraph_for(block)
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            if kind == 'image':
                path = (plan_path.parent / block['path']).resolve()
                width = float(block['width_cm'])
                if width <= 0:
                    raise ValueError('Image width must be positive')
                caption = block.get('caption', '')
                if block.get('evidence_kind') == 'pty_rendered':
                    provenance = (plan_path.parent / block['provenance']).resolve()
                    meta = json.loads(provenance.read_text())
                    if meta.get('kind') != 'pty_rendered' or meta.get('render_sha256') != digest(path):
                        raise ValueError('PTY image does not match its provenance manifest')
                    for field, hash_key in [('source_raw', 'source_sha256'), ('session', 'session_sha256')]:
                        if digest(provenance.parent / meta[field]) != meta[hash_key]:
                            raise ValueError('PTY evidence source changed after rendering')
                    caption += '（真实终端日志渲染，非屏幕截图）'
                run.add_picture(str(path), width=Cm(width))
                end.addprevious(paragraph._p)
                if caption:
                    label = document.add_paragraph(caption, style=block.get('caption_style', 'Caption'))
                    end.addprevious(label._p)
                continue
            text = (plan_path.parent / block['file']).read_text(encoding='utf-8') if 'file' in block else block['text']
            run.text = text
            if kind == 'code':
                apply_font_formatting(
                    run,
                    font_ascii=block.get('font', 'Courier New'),
                    font_eastasia=block.get('font_eastasia'),
                    font_size_pt=block.get('font_size', 9),
                    bold=block.get('bold')
                )
            elif kind == 'paragraph':
                apply_font_formatting(
                    run,
                    font_ascii=block.get('font'),
                    font_eastasia=block.get('font_eastasia'),
                    font_size_pt=block.get('font_size'),
                    bold=block.get('bold')
                )
            end.addprevious(paragraph._p)
    # Delete relationships for removed placeholder images, keeping images used elsewhere.
    for _, part in stories:
        used_ids = {value for node in part.element.iter() for name, value in node.attrib.items()
                    if name in {qn('r:embed'), qn('r:link'), qn('r:id')}}
        for rid, rel in list(part.rels.items()):
            if rel.reltype == RT.IMAGE and rid not in used_ids:
                part.drop_rel(rid)
    _, final_entries, _ = index_document(document)
    all_text = '\n'.join(e['text'] for e in final_entries if e['id'].startswith('P'))
    for token in plan.get('forbid_text', []):
        if token and token in all_text:
            raise ValueError(f'Unresolved placeholder: {token}')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix='.docx', dir=output.parent, delete=False) as f:
        temp = Path(f.name)
    try:
        document.save(temp)
        temp.replace(output)
    finally:
        temp.unlink(missing_ok=True)
    return {'output': str(output.resolve()), 'sha256': digest(output),
            'template_sha256': plan['template_sha256'], 'visual_review': 'pending'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    inspect = commands.add_parser('inspect')
    inspect.add_argument('template', type=Path)
    inspect.add_argument('--out', type=Path, required=True)
    render = commands.add_parser('build')
    render.add_argument('template', type=Path)
    render.add_argument('--plan', type=Path, required=True)
    render.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    try:
        from docx import Document
        if args.command == 'inspect':
            if args.out.exists():
                raise ValueError('Choose a new inspection output path')
            _, entries, _ = index_document(Document(args.template))
            value = {'template_sha256': digest(args.template), 'entries': entries,
                     'note': 'IDs refer only to this template version. Body regions use B IDs; paragraph/cell fields use P/CELL IDs. Inspect layout separately.'}
            write_json(args.out, value)
            print(json.dumps({'inspection': str(args.out), 'entries': len(entries)}))
        else:
            print(json.dumps(build(args.template, args.plan, args.out), ensure_ascii=False, indent=2))
        return 0
    except ImportError:
        print('python-docx is required; use an available document runtime.', file=sys.stderr)
        return 2
    except (OSError, ValueError, KeyError, TypeError, IndexError) as error:
        print(f'Error: {error}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
