"""Behavioral tests for report filling, stable registration and real PTY execution."""
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

SCRIPTS = Path(__file__).resolve().parents[1] / 'killthehw/scripts'
sys.path.insert(0, str(SCRIPTS))
from build_report import build, index_document
from register_requirements import draft, register


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Fixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='killthehw-upgrade-')
        self.root = Path(self.temp.name)
        self.source = self.root / 'assignment.txt'
        self.source.write_text('1. 建表\n1.1. 检查结果\n2. 写报告\n')
        self.job = self.root / 'job'
        self.run_cli('inspect_inputs.py', '--job-dir', self.job, self.source)
        self.run_cli('extract_materials.py', self.job)

    def tearDown(self):
        self.temp.cleanup()

    def run_cli(self, script, *args, expected=0):
        result = subprocess.run([sys.executable, '-B', str(SCRIPTS / script), *map(str, args)],
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return json.loads(result.stdout) if result.stdout else None


class RegistrationTests(Fixture):
    def test_multi_assignment_stable_ids_and_preserved_completion(self):
        job = json.loads((self.job / 'job.json').read_text())
        batch = {'assignment': {'key': 'lab1', 'title': '实验一', 'source': {'input_id': 'I001', 'locator': 'line 1'}},
                 'requirements': [{'key': str(n), 'text': f'要求 {n}', 'source': {'input_id': 'I001', 'locator': f'line {n}'}, 'task_key': 'shared'} for n in (1, 2)]}
        job, ids = register(job, batch)
        self.assertEqual(len(job['tasks']), 1)
        self.assertEqual(len(job['checks']), 2)
        job['tasks'][0]['status'] = 'complete'
        job['checks'][0]['status'] = 'passed'
        same, repeated_ids = register(job, batch)
        self.assertEqual(same, job)
        self.assertEqual(repeated_ids, ids)
        second = copy.deepcopy(batch)
        second['assignment'].update(key='lab2', title='实验二')
        combined, _ = register(job, second)
        self.assertEqual(len(combined['assignments']), 2)
        self.assertEqual(len(combined['tasks']), 2)
        all_ids = [x['id'] for key in ('assignments', 'requirements', 'tasks', 'checks') for x in combined[key]]
        self.assertEqual(len(all_ids), len(set(all_ids)))
        changed = copy.deepcopy(batch)
        changed['requirements'][0]['text'] = '变更后的要求'
        with self.assertRaisesRegex(ValueError, 'changed'):
            register(job, changed)

    def test_numbered_candidates_keep_locations_and_need_review(self):
        self.source.write_text('1. 检查 data[1] 值\n1.1. 检查结果\n2. 写报告\n')
        current = json.loads((self.job / 'job.json').read_text())
        current['inputs'][0]['sha256'] = sha(self.source)
        (self.job / 'job.json').write_text(json.dumps(current))
        self.run_cli('extract_materials.py', self.job)
        out = self.root / 'draft.json'
        self.run_cli('register_requirements.py', 'draft', self.job, '--input', 'I001',
                     '--assignment-key', 'lab1', '--title', '实验一', '--out', out)
        batch = json.loads(out.read_text())
        self.assertEqual([r['number'] for r in batch['requirements']], ['1', '1.1', '2'])
        self.assertEqual(batch['requirements'][1]['source']['locator'], 'line 2')
        self.assertEqual(batch['requirements'][0]['text'], '检查 data[1] 值')
        original = (self.job / 'job.json').read_bytes()
        self.run_cli('register_requirements.py', 'import', self.job, out, '--dry-run')
        self.assertEqual((self.job / 'job.json').read_bytes(), original)
        self.run_cli('register_requirements.py', 'import', self.job, out)
        registered = json.loads((self.job / 'job.json').read_text())
        self.assertTrue(all(c['status'] == 'pending' for c in registered['checks']))
        self.run_cli('register_requirements.py', 'import', self.job, out)
        self.assertEqual(json.loads((self.job / 'job.json').read_text()), registered)


class ReportTests(Fixture):
    def make_template(self):
        from docx import Document
        from docx.shared import Cm
        from PIL import Image
        document = Document()
        document.sections[0].header.paragraphs[0].text = '课程报告'
        cover = document.add_table(rows=2, cols=2)
        cover.cell(0, 0).text = '姓名'
        p = cover.cell(0, 1).paragraphs[0]
        p.add_run('{{姓').bold = True
        p.add_run('名}}').italic = True
        cover.cell(1, 0).merge(cover.cell(1, 1)).text = '课程：待填'
        document.add_paragraph('实验内容', 'Heading 1')
        document.add_paragraph('请在此处填写')
        self.placeholder = self.root / 'placeholder.png'
        self.actual = self.root / 'actual.png'
        Image.new('RGB', (80, 40), 'red').save(self.placeholder)
        Image.new('RGB', (120, 60), 'blue').save(self.actual)
        document.add_picture(str(self.placeholder), width=Cm(3))
        document.add_paragraph('实验内容', 'Heading 1')
        document.add_paragraph('第二实验原文保留')
        template = self.root / 'template.docx'
        document.save(template)
        _, entries, _ = index_document(document)
        start, end = [e['id'] for e in entries if e['id'].startswith('B') and e['text'] == '实验内容']
        cell = next(e['id'] for e in entries if e['id'].startswith('CELL') and e['text'] == '课程：待填')
        plan = {'template_sha256': sha(template),
                'replacements': [{'token': '{{姓名}}', 'value': '测试学生', 'expected_count': 1}],
                'fields': [{'target': cell, 'expect': '课程：待填', 'text': '课程：数据实验'}],
                'sections': [{'start': start, 'end': end, 'expect': '实验内容', 'heading': '实验一结果',
                              'blocks': [{'type': 'paragraph', 'text': '实际结果如下。'},
                                         {'type': 'code', 'text': 'SELECT COUNT(*) FROM scores;'},
                                         {'type': 'image', 'path': 'actual.png', 'width_cm': 8, 'caption': '图 1 实际结果'}]}],
                'forbid_text': ['请在此处填写', '{{姓名}}']}
        return template, plan

    def test_cover_split_runs_duplicate_titles_and_placeholder_images(self):
        from docx import Document
        template, plan = self.make_template()
        original = template.read_bytes()
        plan_path = self.root / 'plan.json'
        plan_path.write_text(json.dumps(plan))
        output = self.root / 'report.docx'
        result = build(template, plan_path, output)
        self.assertEqual(result['visual_review'], 'pending')
        self.assertEqual(template.read_bytes(), original)
        doc = Document(output)
        self.assertEqual(doc.tables[0].cell(0, 1).text, '测试学生')
        self.assertTrue(doc.tables[0].cell(0, 1).paragraphs[0].runs[0].bold)
        self.assertEqual(doc.tables[0].cell(1, 0).text, '课程：数据实验')
        texts = [p.text for p in doc.paragraphs]
        self.assertLess(texts.index('实验一结果'), texts.index('图 1 实际结果'))
        self.assertLess(texts.index('图 1 实际结果'), texts.index('实验内容'))
        self.assertIn('第二实验原文保留', texts)
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, '课程报告')
        with ZipFile(output) as z:
            images = [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
            self.assertNotIn(self.placeholder.read_bytes(), images)
            self.assertIn(self.actual.read_bytes(), images)
        with ZipFile(template) as a, ZipFile(output) as b:
            for name in ('word/styles.xml', 'word/header1.xml'):
                self.assertEqual(a.read(name), b.read(name))

    def test_bad_anchors_or_hash_do_not_write_partial_output(self):
        template, plan = self.make_template()
        output = self.root / 'report.docx'
        plan_path = self.root / 'plan.json'
        for change in ('hash', 'anchor', 'overlap', 'remove_end', 'placeholder'):
            bad = copy.deepcopy(plan)
            if change == 'hash': bad['template_sha256'] = 'stale'
            elif change == 'anchor': bad['sections'][0]['expect'] = 'wrong'
            elif change == 'overlap': bad['sections'].append(copy.deepcopy(bad['sections'][0]))
            elif change == 'remove_end': bad['remove_blocks'] = [bad['sections'][0]['end']]
            else: bad['forbid_text'].append('第二实验原文保留')
            plan_path.write_text(json.dumps(bad))
            with self.assertRaises(ValueError): build(template, plan_path, output)
            self.assertFalse(output.exists())

    def test_font_formatting_and_alignment(self):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        template, plan = self.make_template()
        plan['sections'][0]['blocks'].append({
            'type': 'paragraph',
            'text': '中英混排测试 English',
            'font': 'Times New Roman',
            'font_eastasia': '宋体',
            'font_size': 12,
            'line_spacing': 1.5,
            'alignment': 'center',
            'bold': True
        })
        plan_path = self.root / 'font_plan.json'
        plan_path.write_text(json.dumps(plan))
        output = self.root / 'font_report.docx'
        build(template, plan_path, output)
        self.assertTrue(output.exists())
        doc = Document(output)
        p = next(p for p in doc.paragraphs if '中英混排测试' in p.text)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertAlmostEqual(p.paragraph_format.line_spacing, 1.5)
        run = p.runs[0]
        self.assertTrue(run.bold)
        rFonts = run._r.rPr.find(qn('w:rFonts'))
        self.assertIsNotNone(rFonts)
        self.assertEqual(rFonts.get(qn('w:ascii')), 'Times New Roman')
        self.assertEqual(rFonts.get(qn('w:eastAsia')), '宋体')



class PtyTests(Fixture):
    def test_explicit_stop_preserves_partial_output(self):
        self.run_cli('pty_session.py', 'start', self.job, 'stop-me', '--cwd', self.root,
                     '--', sys.executable, '-c', "import time; print('RUNNING',flush=True); time.sleep(30)")
        self.run_cli('pty_session.py', 'status', self.job, 'stop-me', '--wait', 0.3, expected=1)
        self.run_cli('pty_session.py', 'stop', self.job, 'stop-me')
        state = self.run_cli('pty_session.py', 'status', self.job, 'stop-me', '--wait', 5, expected=1)
        self.assertEqual(state['state'], 'stopped')
        self.assertNotEqual(state['exit_code'], 0)
        self.assertIn(b'RUNNING', (self.job / 'evidence/pty/stop-me/raw.log').read_bytes())

    def test_pty_status_real_output_nonzero_exit_and_render(self):
        self.run_cli('pty_session.py', 'start', self.job, 'one', '--cwd', self.root,
                     '--', sys.executable, '-c', "import os; print('isatty='+str(os.isatty(1))); print('\\x1b[31mACTUAL 4\\x1b[0m')")
        state = self.run_cli('pty_session.py', 'status', self.job, 'one', '--wait', 5)
        session = self.job / 'evidence/pty/one'
        raw = (session / 'raw.log').read_bytes()
        self.assertIn(b'isatty=True', raw)
        self.assertIn(b'ACTUAL 4', raw)
        self.assertEqual(state['raw_sha256'], sha(session / 'raw.log'))
        self.run_cli('pty_session.py', 'start', self.job, 'failed', '--cwd', self.root,
                     '--', sys.executable, '-c', "print('actual error'); raise SystemExit(7)")
        failed = self.run_cli('pty_session.py', 'status', self.job, 'failed', '--wait', 5, expected=1)
        self.assertEqual(failed['exit_code'], 7)
        font = Path('/System/Library/Fonts/Menlo.ttc')
        if font.exists():
            output = self.root / 'render'
            result = self.run_cli('render_terminal.py', session, '--out-dir', output, '--font', font)
            manifest = json.loads((output / 'page-1.json').read_text())
            self.assertEqual(result['kind'], 'pty_rendered')
            self.assertEqual((output / 'source.raw').read_bytes(), raw)
            self.assertEqual(manifest['render_sha256'], sha(output / 'page-1.png'))
            (session / 'raw.log').write_bytes(raw + b'fabricated')
            self.run_cli('render_terminal.py', session, '--out-dir', self.root / 'bad-render', '--font', font, expected=2)

    def test_interactive_wait_is_not_success_and_input_isolated(self):
        self.run_cli('pty_session.py', 'start', self.job, 'interactive', '--cwd', self.root,
                     '--', sys.executable, '-c', "print('READY',flush=True); v=input(); print('GOT:'+v)")
        state = self.run_cli('pty_session.py', 'status', self.job, 'interactive', '--wait', 0.2, expected=1)
        self.assertIn(state['state'], {'starting', 'running'})
        response = self.root / 'response.txt'
        response.write_text('client-response\n')
        self.run_cli('pty_session.py', 'send', self.job, 'interactive', '--file', response)
        self.run_cli('pty_session.py', 'status', self.job, 'interactive', '--wait', 5)
        self.assertIn(b'GOT:client-response', (self.job / 'evidence/pty/interactive/raw.log').read_bytes())
        self.run_cli('pty_session.py', 'send', self.job, 'interactive', '--file', response, expected=2)
        self.run_cli('pty_session.py', 'start', self.job, 'interactive', '--cwd', self.root, '--', 'true', expected=2)


class EnvAndAuditTests(Fixture):

    def test_check_env_runs_successfully(self):
        res = subprocess.run([sys.executable, '-B', str(SCRIPTS / 'check_env.py')],
                             capture_output=True, text=True)
        self.assertEqual(res.returncode, 0)
        self.assertIn('KillTheHW 环境与运行能力体检', res.stdout)
        self.assertIn('模式 1: 整理与排版', res.stdout)

    def test_check_deliverables_summary_flag(self):
        # Initial empty job should produce a summary showing issues
        res = subprocess.run([sys.executable, '-B', str(SCRIPTS / 'check_deliverables.py'),
                              str(self.job), '--summary'],
                             capture_output=True, text=True)
        self.assertIn('KillTheHW 实验报告交付预检审计单', res.stdout)
        self.assertIn('审计结论', res.stdout)


if __name__ == '__main__':
    unittest.main(verbosity=2)

