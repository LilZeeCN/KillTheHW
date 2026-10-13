#!/usr/bin/env python3
"""Generate sample template.docx and evidence image for quickstart."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont


def set_run_fonts(run, ascii_font="Times New Roman", eastasia_font="宋体"):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    run.font.name = ascii_font
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), eastasia_font)


def create_evidence_image(path):
    img = Image.new("RGB", (640, 240), color=(30, 30, 30))
    draw = ImageDraw.Draw(img)
    # Draw simple terminal representation
    draw.rectangle([0, 0, 640, 30], fill=(50, 50, 50))
    draw.ellipse([10, 10, 20, 20], fill=(255, 95, 86))
    draw.ellipse([26, 10, 36, 20], fill=(255, 189, 46))
    draw.ellipse([42, 10, 52, 20], fill=(39, 201, 63))
    
    text = (
        "student@lab-vm:~$ uname -a\n"
        "Linux lab-vm 6.6.0-generic #1 SMP PREEMPT_DYNAMIC x86_64 GNU/Linux\n"
        "student@lab-vm:~$ df -h /\n"
        "Filesystem      Size  Used Avail Use% Mounted on\n"
        "/dev/sda1        50G   14G   34G  30% /\n"
        "student@lab-vm:~$ ping -c 2 127.0.0.1\n"
        "64 bytes from 127.0.0.1: icmp_seq=1 ttl=64 time=0.038 ms\n"
        "2 packets transmitted, 2 received, 0% packet loss"
    )
    draw.text((15, 45), text, fill=(0, 255, 128))
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def create_template(path):
    doc = Document()
    
    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("计算机科学与技术学院 实验报告")
    run_title.bold = True
    run_title.font.size = Pt(18)
    set_run_fonts(run_title, "SimHei", "黑体")
    
    # Cover Table
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    fields = [
        ("课程名称：", "{{课程名称}}"),
        ("学生姓名：", "{{学生姓名}}"),
        ("学号信息：", "{{学号}}"),
    ]
    for row_idx, (label, val) in enumerate(fields):
        r_label = table.cell(row_idx, 0).paragraphs[0].add_run(label)
        set_run_fonts(r_label, "SimSun", "宋体")
        r_val = table.cell(row_idx, 1).paragraphs[0].add_run(val)
        set_run_fonts(r_val, "SimSun", "宋体")
        
    doc.add_paragraph() # spacer
    
    # Section 1
    h1 = doc.add_paragraph("一、实验内容与步骤", "Heading 1")
    p1 = doc.add_paragraph("请在此处填写实验步骤及运行结果截图。")
    
    # Section 2
    h2 = doc.add_paragraph("二、思考题与总结", "Heading 1")
    p2 = doc.add_paragraph("请在此处回答指导书思考题及心得。")
    
    # End anchor
    doc.add_paragraph("实验报告结束。")
    
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    base = Path(__file__).parent
    create_template(base / "template.docx")
    create_evidence_image(base / "evidence" / "term_uname.png")
    print("Quickstart assets generated successfully.")
